"""
Real files of every supported format, built in memory for the upload tests.

These are genuine PDFs, Word documents and spreadsheets rather than fixtures
checked in as binaries, for two reasons: a committed binary is a thing nobody
can review in a diff, and the point of these tests is that the *parsers* work
-- so the bytes have to be produced the way a real file is, not replayed from
a blob that happened to pass once.

The PDF is assembled by hand because nothing in the dependency list writes one
(pypdf reads them; reportlab is not installed). It is small but structurally
complete -- object offsets, xref table, trailer and EOF marker -- because a PDF
missing any of those is rejected by pypdf before a single character of text is
read, which is a test of nothing.
"""

import io
import json


def pdf_bytes(lines) -> bytes:
    """A one-page PDF whose text layer is `lines`, with a valid xref."""
    text_operations = "\n".join(
        f"BT /F1 12 Tf 72 {720 - index * 18} Td ({line}) Tj ET" for index, line in enumerate(lines)
    )
    stream = text_operations.encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    return bytes(out)


def docx_bytes(heading, paragraphs) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading(heading, 1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(sheet_name, rows) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def csv_bytes(rows) -> bytes:
    return "\n".join(",".join(str(cell) for cell in row) for row in rows).encode("utf-8")


def json_bytes(payload) -> bytes:
    return json.dumps(payload).encode("utf-8")
