"""
Tests for the knowledge-base feature: extraction, the processing pipeline,
retrieval isolation and confidence, and the answering behaviour that hangs off
them.

Kept separate from tests.py, which covers the agent itself, because these are
about a different thing going wrong: an agent that answers badly is a prompt
problem, while a document that indexed as gibberish or a tenant that can read
another tenant's price list is a correctness problem.
"""

import io
import os
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from rest_framework.test import APIClient

from apps.contacts.models import Contact
from apps.inbox.models import Conversation
from apps.messaging.models import Message
from apps.tenants.models import Tenant
from apps.whatsapp.models import WhatsAppAccount
from apps.whatsapp.providers.base import SendResult

from . import extraction, knowledge, processing, services
from .models import (
    AIAgentConfig,
    AIConversationState,
    KnowledgeAnswer,
    KnowledgeBase,
    KnowledgeChunk,
    KnowledgeDocument,
)
from .providers.anthropic import AIResponse

PRICING = """Premium Plan Pricing

The premium plan costs 4999 rupees per month and includes unlimited messages.

The starter plan costs 999 rupees per month and includes 1000 messages."""

WARRANTY = """Warranty Terms

Every appliance carries a two year manufacturer warranty from the date of purchase.

Warranty claims require the original invoice."""


# ---------------------------------------------------------------------------
# Fixture builders
# ---------------------------------------------------------------------------

def _escape_pdf(text):
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def make_pdf(pages) -> bytes:
    """
    A minimal but genuinely valid PDF carrying one text block per page.

    Built by hand rather than mocked: the point of the PDF test is that a real
    parser reads a real file, and a mocked PdfReader would pass just as
    happily against an extractor that does not work.
    """
    page_count = len(pages)
    page_ids = [3 + index * 2 for index in range(page_count)]
    content_ids = [4 + index * 2 for index in range(page_count)]
    font_id = 3 + page_count * 2

    objects = [
        (1, "<< /Type /Catalog /Pages 2 0 R >>"),
        (2, f"<< /Type /Pages /Kids [{' '.join(f'{pid} 0 R' for pid in page_ids)}] /Count {page_count} >>"),
    ]
    for index, text in enumerate(pages):
        page_id, content_id = page_ids[index], content_ids[index]
        objects.append((
            page_id,
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>",
        ))
        body = "BT /F1 12 Tf 72 720 Td 14 TL\n" + "\n".join(
            f"({_escape_pdf(line)}) Tj T*" for line in text.split("\n")
        ) + "\nET"
        objects.append((content_id, f"<< /Length {len(body)} >>\nstream\n{body}\nendstream"))
    objects.append((font_id, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"))

    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for number, payload in sorted(objects):
        offsets[number] = len(out)
        out += f"{number} 0 obj\n{payload}\nendobj\n".encode("latin-1")

    xref_at = len(out)
    highest = max(offsets)
    out += f"xref\n0 {highest + 1}\n".encode("latin-1")
    out += b"0000000000 65535 f \n"
    for number in range(1, highest + 1):
        out += f"{offsets.get(number, 0):010d} 00000 n \n".encode("latin-1")
    out += f"trailer\n<< /Size {highest + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode("latin-1")
    return bytes(out)


def make_docx(paragraphs, headings=()) -> bytes:
    import docx

    document = docx.Document()
    for heading in headings:
        document.add_heading(heading, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(rows, sheet_name="Prices") -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _account(tenant, user, name="Line", pid="1"):
    account = WhatsAppAccount(
        tenant=tenant, user=user, name=name, phone_number_id=pid, status=WhatsAppAccount.Status.CONNECTED,
    )
    account.set_access_token("token")
    account.set_app_secret("secret")
    account.save()
    return account


class KnowledgeFixture(TestCase):
    def setUp(self):
        self.tenant = Tenant.objects.create(
            name="Acme", slug=f"kb-{self._testMethodName}"[:50], kind=Tenant.Kind.CUSTOMER,
        )
        self.owner = get_user_model().objects.create_user(
            username=f"kb-owner-{self._testMethodName}"[:140], password="x", tenant=self.tenant, credits=500,
        )
        self.account = _account(self.tenant, self.owner)
        self.contact = Contact.objects.create(
            tenant=self.tenant, created_by=self.owner, name="Priya", phone="+15557770001",
        )
        self.conversation = Conversation.objects.create(
            tenant=self.tenant, whatsapp_account=self.account, contact=self.contact,
        )
        self.base = KnowledgeBase.objects.create(tenant=self.tenant, name="Main", is_default=True)
        self.client = APIClient()
        self.client.force_authenticate(self.owner)

    def _config(self, **kwargs):
        config = AIAgentConfig.objects.create(
            tenant=self.tenant, is_enabled=True, knowledge_base=self.base, **kwargs,
        )
        config.set_api_key("sk-test")
        config.save()
        return config

    def _document(self, content=PRICING, title="Pricing", base=None, tenant=None):
        document = KnowledgeDocument.objects.create(
            tenant=tenant or self.tenant, knowledge_base=base or self.base, title=title,
            content=content, status=KnowledgeDocument.Status.READY,
        )
        document.chunk_count = knowledge.rebuild(document)
        document.save(update_fields=["chunk_count"])
        return document

    def _inbound(self, body="How much is the premium plan?"):
        return Message.objects.create(
            tenant=self.tenant, conversation=self.conversation, whatsapp_account=self.account,
            contact=self.contact, to_phone=self.contact.phone, direction=Message.Direction.INBOUND,
            body=body, status=Message.Status.DELIVERED,
        )


def _reply(text):
    return AIResponse(text=text, input_tokens=100, output_tokens=20, latency_ms=120)


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------

class ExtractionValidationTests(TestCase):
    def test_a_supported_extension_passes(self):
        self.assertEqual(extraction.validate_upload("prices.pdf", 1000, "application/pdf"), ".pdf")

    def test_an_unsupported_extension_is_refused_by_name(self):
        with self.assertRaises(extraction.ExtractionError) as caught:
            extraction.validate_upload("virus.exe", 1000, "application/octet-stream")
        self.assertIn(".exe", str(caught.exception))

    def test_a_file_with_no_extension_is_refused(self):
        with self.assertRaises(extraction.ExtractionError):
            extraction.validate_upload("README", 1000, "text/plain")

    def test_an_empty_file_is_refused(self):
        with self.assertRaises(extraction.ExtractionError):
            extraction.validate_upload("prices.pdf", 0, "application/pdf")

    def test_an_oversized_file_is_refused(self):
        with self.assertRaises(extraction.ExtractionError) as caught:
            extraction.validate_upload("huge.pdf", extraction.MAX_FILE_BYTES + 1, "application/pdf")
        self.assertIn("MB", str(caught.exception))

    def test_a_wrong_content_type_does_not_block_a_valid_extension(self):
        """Phones send nonsense content types; the extension is authoritative."""
        self.assertEqual(extraction.validate_upload("data.csv", 500, "application/octet-stream"), ".csv")


class PdfExtractionTests(TestCase):
    def test_text_and_page_numbers_come_out(self):
        pdf = make_pdf(["Premium plan costs 4999 rupees.", "Warranty lasts two years."])

        sections = extraction.extract(io.BytesIO(pdf), "prices.pdf")

        self.assertEqual(len(sections), 2)
        self.assertEqual(sections[0].page, 1)
        self.assertEqual(sections[1].page, 2)
        self.assertIn("4999", sections[0].text)
        self.assertIn("two years", sections[1].text)

    def test_a_pdf_with_no_text_layer_says_so(self):
        """A scan is images, not text — the user needs telling, not an empty index."""
        pdf = make_pdf([""])
        with self.assertRaises(extraction.ExtractionError) as caught:
            extraction.extract(io.BytesIO(pdf), "scan.pdf")
        self.assertIn("scan", str(caught.exception).lower())

    def test_a_corrupt_pdf_fails_readably_rather_than_crashing(self):
        with self.assertRaises(extraction.ExtractionError):
            extraction.extract(io.BytesIO(b"not a pdf at all"), "broken.pdf")


class DocxExtractionTests(TestCase):
    def test_paragraphs_and_headings_come_out(self):
        blob = make_docx(["The premium plan costs 4999 rupees per month."], headings=["Pricing"])

        sections = extraction.extract(io.BytesIO(blob), "pricing.docx")

        self.assertTrue(sections)
        self.assertIn("4999", " ".join(section.text for section in sections))
        self.assertIn("Pricing", [section.heading for section in sections])

    def test_a_corrupt_docx_fails_readably(self):
        with self.assertRaises(extraction.ExtractionError):
            extraction.extract(io.BytesIO(b"nonsense"), "broken.docx")


class TextExtractionTests(TestCase):
    def test_plain_text_is_read(self):
        sections = extraction.extract(io.BytesIO(PRICING.encode()), "faq.txt")
        self.assertIn("4999", " ".join(section.text for section in sections))

    def test_markdown_headings_become_section_headings(self):
        body = "# Refunds\n\nWe refund within 14 days.\n\n# Delivery\n\nDelivery takes two days."
        sections = extraction.extract(io.BytesIO(body.encode()), "policy.md")
        headings = [section.heading for section in sections]
        self.assertIn("Refunds", headings)
        self.assertIn("Delivery", headings)

    def test_latin1_text_does_not_blow_up(self):
        sections = extraction.extract(io.BytesIO("Café prices".encode("latin-1")), "menu.txt")
        self.assertTrue(sections)


class TabularExtractionTests(TestCase):
    def test_csv_rows_become_findable_sentences(self):
        """
        A raw row shares no words with the question. Pairing each cell with its
        header is what makes "how much is the premium plan" match at all.
        """
        csv_bytes = b"Plan,Price,Messages\nPremium,4999,Unlimited\nStarter,999,1000\n"

        sections = extraction.extract(io.BytesIO(csv_bytes), "plans.csv")

        text = " ".join(section.text for section in sections)
        self.assertIn("Plan: Premium", text)
        self.assertIn("Price: 4999", text)

    def test_a_csv_with_no_header_still_reads(self):
        sections = extraction.extract(io.BytesIO(b"Premium,4999\nStarter,999\n"), "plans.csv")
        self.assertTrue(sections)

    def test_xlsx_rows_are_read_with_the_sheet_name(self):
        blob = make_xlsx([["Plan", "Price"], ["Premium", 4999], ["Starter", 999]], sheet_name="Rates")

        sections = extraction.extract(io.BytesIO(blob), "plans.xlsx")

        self.assertTrue(sections)
        self.assertEqual(sections[0].heading, "Rates")
        self.assertIn("4999", sections[0].text)

    def test_an_empty_csv_is_refused(self):
        with self.assertRaises(extraction.ExtractionError):
            extraction.extract(io.BytesIO(b""), "empty.csv")


class JsonExtractionTests(TestCase):
    """
    JSON is flattened into labelled sentences for the same reason CSV rows are:
    a question shares words with "refund window days: 30" and shares none with
    a wall of braces and quotes.
    """

    def test_nested_keys_become_labelled_sentences(self):
        blob = b'{"refund_policy": {"window_days": 30, "condition": "unused"}}'

        sections = extraction.extract(io.BytesIO(blob), "policies.json")

        text = " ".join(section.text for section in sections)
        self.assertIn("refund policy window days: 30", text)
        self.assertIn("refund policy condition: unused", text)

    def test_a_top_level_array_becomes_one_section_per_item(self):
        blob = b'[{"name": "Premium", "price": 4999}, {"name": "Starter", "price": 999}]'

        sections = extraction.extract(io.BytesIO(blob), "plans.json")

        self.assertEqual(len(sections), 2)
        self.assertIn("4999", sections[0].text)
        self.assertIn("999", sections[1].text)

    def test_a_list_of_scalars_reads_as_one_sentence(self):
        blob = b'{"sizes": ["small", "medium", "large"]}'

        sections = extraction.extract(io.BytesIO(blob), "sizes.json")

        self.assertIn("sizes: small, medium, large", sections[0].text)

    def test_each_top_level_key_gets_its_own_section(self):
        # So a retrieval can return the shipping terms without the returns
        # policy riding along with them.
        blob = b'{"shipping": {"days": 2}, "returns": {"days": 14}}'

        sections = extraction.extract(io.BytesIO(blob), "terms.json")

        self.assertEqual([section.heading for section in sections], ["shipping", "returns"])

    def test_empty_values_are_dropped_rather_than_labelled(self):
        blob = b'{"product": {"name": "Widget", "discontinued": null, "note": ""}}'

        sections = extraction.extract(io.BytesIO(blob), "product.json")

        text = " ".join(section.text for section in sections)
        self.assertIn("Widget", text)
        self.assertNotIn("discontinued", text)
        self.assertNotIn("note", text)

    def test_malformed_json_is_refused_with_a_reason(self):
        with self.assertRaises(extraction.ExtractionError) as caught:
            extraction.extract(io.BytesIO(b'{"unclosed": '), "broken.json")

        self.assertIn("not valid JSON", str(caught.exception))

    def test_json_is_an_accepted_upload(self):
        self.assertEqual(extraction.validate_upload("policies.json", 500, "application/json"), ".json")


class CleaningTests(TestCase):
    def test_runs_of_whitespace_collapse_but_paragraphs_survive(self):
        cleaned = extraction.clean("A   line\n\n\n\nAnother   line")
        self.assertEqual(cleaned, "A line\n\nAnother line")

    def test_null_bytes_are_stripped(self):
        self.assertNotIn("\x00", extraction.clean("bad\x00text"))


# ---------------------------------------------------------------------------
# Processing pipeline
# ---------------------------------------------------------------------------

class ProcessingTests(KnowledgeFixture):
    def _upload(self, content, filename, content_type="application/pdf"):
        return KnowledgeDocument.objects.create(
            tenant=self.tenant, knowledge_base=self.base, title=filename,
            source_type=KnowledgeDocument.Source.FILE,
            file=SimpleUploadedFile(filename, content, content_type=content_type),
            original_filename=filename, size_bytes=len(content),
            status=KnowledgeDocument.Status.PENDING,
        )

    def test_a_pdf_becomes_ready_with_chunks_and_pages(self):
        document = self._upload(make_pdf(["Premium plan costs 4999 rupees.", "Warranty lasts two years."]), "prices.pdf")

        processing.process_document(document.pk)

        document.refresh_from_db()
        self.assertEqual(document.status, KnowledgeDocument.Status.READY)
        self.assertGreater(document.chunk_count, 0)
        self.assertEqual(document.page_count, 2)
        self.assertIsNotNone(document.indexed_at)
        self.assertEqual(document.error, "")
        self.assertIn("4999", document.content)

    def test_chunks_carry_the_page_they_came_from(self):
        document = self._upload(make_pdf(["Page one text about pricing.", "Page two text about warranty."]), "doc.pdf")

        processing.process_document(document.pk)

        pages = set(KnowledgeChunk.objects.filter(document=document).values_list("page", flat=True))
        self.assertEqual(pages, {1, 2})

    def test_a_citation_names_the_file_and_page(self):
        document = self._upload(make_pdf(["Premium plan costs 4999 rupees."]), "Pricing.pdf")
        processing.process_document(document.pk)

        chunk = KnowledgeChunk.objects.filter(document=document).first()

        self.assertEqual(chunk.citation, "Pricing.pdf, page 1")

    def test_an_unreadable_file_fails_with_a_reason_a_user_can_act_on(self):
        document = self._upload(b"this is not a pdf", "broken.pdf")

        processing.process_document(document.pk)

        document.refresh_from_db()
        self.assertEqual(document.status, KnowledgeDocument.Status.FAILED)
        self.assertTrue(document.error)
        self.assertEqual(document.chunk_count, 0)

    def test_a_failed_document_is_never_retrieved_from(self):
        document = self._upload(b"garbage", "broken.pdf")
        processing.process_document(document.pk)

        # Force a chunk in, to prove the status filter is what excludes it
        # rather than merely there being nothing to find.
        KnowledgeChunk.objects.create(
            document=document, tenant=self.tenant, knowledge_base=self.base, ordinal=0,
            content="The premium plan costs 4999 rupees.",
        )

        retrieval = knowledge.retrieve(self.tenant.pk, "premium plan price", knowledge_base_id=self.base.pk)
        self.assertEqual(retrieval.hits, [])

    def test_reprocessing_replaces_chunks_rather_than_doubling_them(self):
        document = self._upload(make_pdf(["Premium plan costs 4999 rupees."]), "prices.pdf")
        processing.process_document(document.pk)
        first = KnowledgeChunk.objects.filter(document=document).count()

        processing.process_document(document.pk)

        self.assertEqual(KnowledgeChunk.objects.filter(document=document).count(), first)

    def test_deleting_a_document_removes_its_chunks(self):
        document = self._upload(make_pdf(["Premium plan costs 4999 rupees."]), "prices.pdf")
        processing.process_document(document.pk)
        self.assertTrue(KnowledgeChunk.objects.filter(document=document).exists())

        document.delete()

        self.assertFalse(KnowledgeChunk.objects.filter(document_id=document.pk).exists())

    def test_a_file_missing_from_storage_says_to_upload_it_again(self):
        # A row whose bytes are gone -- a media volume that lost its contents,
        # or an API server that does not share storage with the worker. The
        # uploader needs to be told the one thing that fixes it, not the name
        # of the exception class.
        document = self._upload(make_pdf(["Premium plan costs 4999 rupees."]), "prices.pdf")
        os.remove(document.file.path)

        processing.process_document(document.pk)

        document.refresh_from_db()
        self.assertEqual(document.status, KnowledgeDocument.Status.FAILED)
        self.assertIn("upload it again", document.error.lower())
        self.assertNotIn("FileNotFoundError", document.error)

    def test_deleting_a_document_removes_the_stored_file(self):
        document = self._upload(make_pdf(["Premium plan costs 4999 rupees."]), "prices.pdf")
        path = document.file.path
        self.assertTrue(os.path.exists(path))

        document.delete()

        self.assertFalse(os.path.exists(path))

    def test_deleting_a_knowledge_base_removes_its_documents_files(self):
        # The cascade is the case that matters: it is also the path a tenant
        # deletion takes, and a fast delete would skip the file entirely.
        document = self._upload(make_pdf(["Premium plan costs 4999 rupees."]), "prices.pdf")
        path = document.file.path

        self.base.delete()

        self.assertFalse(os.path.exists(path))

    def test_a_document_whose_file_is_already_gone_still_deletes(self):
        document = self._upload(make_pdf(["Premium plan costs 4999 rupees."]), "prices.pdf")
        os.remove(document.file.path)

        document.delete()

        self.assertFalse(KnowledgeDocument.objects.filter(pk=document.pk).exists())

    def test_a_typed_document_is_indexed_without_a_file(self):
        document = KnowledgeDocument.objects.create(
            tenant=self.tenant, knowledge_base=self.base, title="FAQ", content=PRICING,
        )

        processing.process_document(document.pk)

        document.refresh_from_db()
        self.assertEqual(document.status, KnowledgeDocument.Status.READY)
        self.assertGreater(document.chunk_count, 0)


# ---------------------------------------------------------------------------
# Retrieval: isolation and confidence
# ---------------------------------------------------------------------------

class RetrievalIsolationTests(KnowledgeFixture):
    def test_another_tenants_documents_are_never_retrieved(self):
        other = Tenant.objects.create(name="Rival", slug="kb-rival", kind=Tenant.Kind.CUSTOMER)
        other_user = get_user_model().objects.create_user(username="kb-rival-owner", password="x", tenant=other)
        other_base = KnowledgeBase.objects.create(tenant=other, name="Theirs")
        secret = KnowledgeDocument.objects.create(
            tenant=other, knowledge_base=other_base, created_by=other_user,
            title="Secret pricing", content="Our premium plan costs 1 rupee.",
            status=KnowledgeDocument.Status.READY,
        )
        knowledge.rebuild(secret)

        retrieval = knowledge.retrieve(self.tenant.pk, "premium plan price")

        self.assertEqual(retrieval.hits, [])

    def test_a_named_base_never_reaches_another_base(self):
        summer = KnowledgeBase.objects.create(tenant=self.tenant, name="Summer")
        self._document(content="The summer special costs 199 rupees.", title="Summer", base=summer)
        self._document(content=PRICING, title="Pricing", base=self.base)

        retrieval = knowledge.retrieve(self.tenant.pk, "summer special price", knowledge_base_id=self.base.pk)

        titles = [source["title"] for source in retrieval.sources]
        self.assertNotIn("Summer", titles)

    def test_naming_no_base_searches_every_base_the_tenant_has(self):
        summer = KnowledgeBase.objects.create(tenant=self.tenant, name="Summer")
        self._document(content="The summer special costs 199 rupees.", title="Summer", base=summer)

        retrieval = knowledge.retrieve(self.tenant.pk, "summer special")

        self.assertTrue(retrieval.hits)

    def test_an_inactive_base_is_excluded_when_named(self):
        self._document()
        KnowledgeBase.objects.filter(pk=self.base.pk).update(is_active=False)

        retrieval = knowledge.retrieve(self.tenant.pk, "premium plan", knowledge_base_id=self.base.pk)

        self.assertEqual(retrieval.hits, [])

    def test_an_inactive_document_is_excluded(self):
        document = self._document()
        KnowledgeDocument.objects.filter(pk=document.pk).update(is_active=False)

        self.assertEqual(knowledge.retrieve(self.tenant.pk, "premium plan").hits, [])


class ConfidenceTests(KnowledgeFixture):
    def test_a_squarely_answered_question_is_a_good_match(self):
        self._document()

        retrieval = knowledge.retrieve(self.tenant.pk, "premium plan pricing")

        self.assertEqual(retrieval.confidence, KnowledgeAnswer.Confidence.GOOD)
        self.assertGreaterEqual(retrieval.relevance, knowledge.GOOD_RELEVANCE)

    def test_an_unrelated_question_finds_nothing(self):
        self._document()

        retrieval = knowledge.retrieve(self.tenant.pk, "do you sell helicopters")

        self.assertEqual(retrieval.confidence, KnowledgeAnswer.Confidence.NONE)
        self.assertEqual(retrieval.relevance, 0.0)

    def test_an_empty_knowledge_base_finds_nothing(self):
        retrieval = knowledge.retrieve(self.tenant.pk, "premium plan")

        self.assertFalse(retrieval)
        self.assertEqual(retrieval.confidence, KnowledgeAnswer.Confidence.NONE)

    def test_a_question_of_only_stop_words_finds_nothing(self):
        self._document()
        self.assertEqual(knowledge.retrieve(self.tenant.pk, "what is the").hits, [])

    def test_relevance_never_exceeds_one(self):
        self._document()
        self.assertLessEqual(knowledge.retrieve(self.tenant.pk, "premium plan pricing rupees").relevance, 1.0)

    def test_sources_are_deduplicated_per_document(self):
        self._document(content=PRICING + "\n\nThe premium plan also includes support.", title="Pricing")

        retrieval = knowledge.retrieve(self.tenant.pk, "premium plan", limit=5)

        self.assertEqual(len({source["document_id"] for source in retrieval.sources}), len(retrieval.sources))


# ---------------------------------------------------------------------------
# Answering
# ---------------------------------------------------------------------------

@patch("apps.messaging.services.ConnectionManager.send_text_message")
class KnowledgeAnswerTests(KnowledgeFixture):
    def test_a_covered_question_is_answered_and_recorded_with_its_sources(self, mock_send):
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.1")
        self._config()
        self._document()
        inbound = self._inbound()

        with patch.object(services, "complete", return_value=_reply("The premium plan is 4999 rupees a month.")):
            self.assertTrue(services.answer(self.conversation, inbound))

        record = KnowledgeAnswer.objects.get(is_test=False)
        self.assertTrue(record.answered)
        self.assertEqual(record.confidence, KnowledgeAnswer.Confidence.GOOD)
        self.assertEqual([source["title"] for source in record.sources], ["Pricing"])
        self.assertIsNotNone(record.message_id)

    def test_an_uncovered_question_never_reaches_the_model(self, mock_send):
        """
        The cost-control and anti-hallucination path are the same path: with
        nothing relevant retrieved there is nothing to answer from, so no
        request is made and nothing can be invented.
        """
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.2")
        self._config(fallback_action=AIAgentConfig.Fallback.SAY_UNKNOWN)
        self._document()
        inbound = self._inbound("Do you sell helicopters?")

        with patch.object(services, "complete") as mock_model:
            services.answer(self.conversation, inbound)

        mock_model.assert_not_called()
        record = KnowledgeAnswer.objects.get()
        self.assertFalse(record.answered)
        self.assertEqual(record.confidence, KnowledgeAnswer.Confidence.NONE)
        self.assertEqual(record.fallback_used, AIAgentConfig.Fallback.SAY_UNKNOWN)

    def test_the_say_unknown_fallback_admits_it_rather_than_guessing(self, mock_send):
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.3")
        self._config(fallback_action=AIAgentConfig.Fallback.SAY_UNKNOWN)
        self._document()

        services.answer(self.conversation, self._inbound("Do you sell helicopters?"))

        sent = Message.objects.filter(direction=Message.Direction.OUTBOUND).latest("pk")
        self.assertIn("don't have that information", sent.body)

    def test_the_handoff_fallback_hands_over_instead_of_answering(self, mock_send):
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.4")
        self._config(fallback_action=AIAgentConfig.Fallback.HANDOFF)
        self._document()

        services.answer(self.conversation, self._inbound("Do you sell helicopters?"))

        state = AIConversationState.objects.get(conversation=self.conversation)
        self.assertIsNotNone(state.handed_off_at)

    def test_the_notify_team_fallback_notifies_and_replies(self, mock_send):
        from apps.notifications.models import Notification
        from apps.teams.models import TeamMember

        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.5")
        TeamMember.objects.create(tenant=self.tenant, user=self.owner, role=TeamMember.Role.AGENT)
        self._config(fallback_action=AIAgentConfig.Fallback.NOTIFY_TEAM)
        self._document()

        services.answer(self.conversation, self._inbound("Do you sell helicopters?"))

        self.assertTrue(Notification.objects.filter(tenant=self.tenant).exists())

    def test_asking_for_a_person_hands_over_without_a_model_call(self, mock_send):
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.6")
        self._config()
        self._document()

        with patch.object(services, "complete") as mock_model:
            services.answer(self.conversation, self._inbound("I want to speak to someone please"))

        mock_model.assert_not_called()
        self.assertIsNotNone(AIConversationState.objects.get(conversation=self.conversation).handed_off_at)

    def test_the_model_refusing_still_produces_the_fallback(self, mock_send):
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.7")
        self._config(fallback_action=AIAgentConfig.Fallback.SAY_UNKNOWN)
        self._document()

        with patch.object(services, "complete", return_value=_reply(services.HANDOFF_TOKEN)):
            services.answer(self.conversation, self._inbound())

        sent = Message.objects.filter(direction=Message.Direction.OUTBOUND).latest("pk")
        self.assertIn("don't have that information", sent.body)

    def test_only_the_configured_knowledge_base_is_used(self, mock_send):
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.8")
        summer = KnowledgeBase.objects.create(tenant=self.tenant, name="Summer")
        self._document(content="The summer special costs 199 rupees.", title="Summer", base=summer)
        self._config()  # points at self.base, which has nothing in it

        with patch.object(services, "complete") as mock_model:
            services.answer(self.conversation, self._inbound("How much is the summer special?"))

        mock_model.assert_not_called()

    def test_a_human_owning_the_chat_stops_the_agent(self, mock_send):
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.9")
        self._config(reply_scope=AIAgentConfig.ReplyScope.UNASSIGNED)
        self._document()
        Conversation.objects.filter(pk=self.conversation.pk).update(assigned_agent=self.owner)
        self.conversation.refresh_from_db()

        with patch.object(services, "complete") as mock_model:
            self.assertFalse(services.answer(self.conversation, self._inbound()))

        mock_model.assert_not_called()

    def test_every_message_scope_answers_even_an_assigned_chat(self, mock_send):
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.10")
        self._config(reply_scope=AIAgentConfig.ReplyScope.EVERY)
        self._document()
        Conversation.objects.filter(pk=self.conversation.pk).update(assigned_agent=self.owner)
        self.conversation.refresh_from_db()

        with patch.object(services, "complete", return_value=_reply("It is 4999 rupees.")):
            self.assertTrue(services.answer(self.conversation, self._inbound()))

    def test_outside_hours_scope_stays_quiet_while_the_team_is_open(self, mock_send):
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.11")
        self._config(reply_scope=AIAgentConfig.ReplyScope.OUTSIDE_HOURS)
        self._document()

        with patch("apps.teams.services.is_open_now", return_value=True):
            with patch.object(services, "complete") as mock_model:
                self.assertFalse(services.answer(self.conversation, self._inbound()))

        mock_model.assert_not_called()

    def test_outside_hours_scope_answers_when_the_team_is_closed(self, mock_send):
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.12")
        self._config(reply_scope=AIAgentConfig.ReplyScope.OUTSIDE_HOURS)
        self._document()

        with patch("apps.teams.services.is_open_now", return_value=False):
            with patch.object(services, "complete", return_value=_reply("It is 4999 rupees.")):
                self.assertTrue(services.answer(self.conversation, self._inbound()))

    def test_an_already_handed_over_conversation_is_left_alone(self, mock_send):
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.13")
        self._config()
        self._document()
        AIConversationState.objects.create(
            conversation=self.conversation, handed_off_at="2026-01-01T00:00:00Z",
        )

        with patch.object(services, "complete") as mock_model:
            self.assertFalse(services.answer(self.conversation, self._inbound()))

        mock_model.assert_not_called()

    def test_the_reply_goes_out_through_the_normal_send_path(self, mock_send):
        """Never a second sender: limits, quiet hours and consent all still apply."""
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.14")
        self._config()
        self._document()

        with patch.object(services, "complete", return_value=_reply("It is 4999 rupees.")):
            services.answer(self.conversation, self._inbound())

        sent = Message.objects.filter(direction=Message.Direction.OUTBOUND).latest("pk")
        self.assertTrue(sent.is_automated)
        self.assertEqual(sent.kind, Message.Kind.SERVICE)
        self.assertEqual(sent.conversation_id, self.conversation.pk)

    def test_an_unsubscribed_contact_is_never_messaged(self, mock_send):
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.kb.15")
        self._config()
        self._document()
        Contact.objects.filter(pk=self.contact.pk).update(subscribed=False)

        with patch.object(services, "complete", return_value=_reply("It is 4999 rupees.")):
            services.answer(self.conversation, self._inbound())

        sent = Message.objects.filter(direction=Message.Direction.OUTBOUND).latest("pk")
        # The gate runs at send time and marks it SKIPPED rather than sending.
        self.assertIn(sent.status, (Message.Status.SKIPPED, Message.Status.QUEUED))
        if sent.status == Message.Status.SKIPPED:
            self.assertIn("unsub", sent.error_message.lower())


# ---------------------------------------------------------------------------
# CRM integration
# ---------------------------------------------------------------------------

@patch("apps.messaging.services.ConnectionManager.send_text_message")
class KnowledgeCrmTests(KnowledgeFixture):
    def test_the_question_lands_on_the_contacts_timeline(self, mock_send):
        from apps.crm.models import Activity

        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.crm.1")
        self._config(crm_logging_enabled=True)
        self._document()

        with patch.object(services, "complete", return_value=_reply("It is 4999 rupees.")):
            services.answer(self.conversation, self._inbound("What is the price of the premium plan?"))

        entry = Activity.objects.filter(contact=self.contact, kind=Activity.Kind.AUTOMATION).first()
        self.assertIsNotNone(entry)
        self.assertEqual(entry.metadata["topic"], "pricing")

    def test_tagging_is_off_unless_asked_for(self, mock_send):
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.crm.2")
        self._config(crm_logging_enabled=True, crm_tagging_enabled=False)
        self._document()

        with patch.object(services, "complete", return_value=_reply("It is 4999 rupees.")):
            services.answer(self.conversation, self._inbound("What is the price?"))

        self.assertEqual(self.contact.tags.count(), 0)

    def test_tagging_applies_the_topic_tag_when_switched_on(self, mock_send):
        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.crm.3")
        self._config(crm_logging_enabled=True, crm_tagging_enabled=True)
        self._document()

        with patch.object(services, "complete", return_value=_reply("It is 4999 rupees.")):
            services.answer(self.conversation, self._inbound("What is the price of the premium plan?"))

        self.assertIn("Pricing enquiry", [tag.name for tag in self.contact.tags.all()])

    def test_crm_logging_can_be_switched_off_entirely(self, mock_send):
        from apps.crm.models import Activity

        mock_send.return_value = SendResult(success=True, provider_message_id="wamid.crm.4")
        self._config(crm_logging_enabled=False)
        self._document()

        with patch.object(services, "complete", return_value=_reply("It is 4999 rupees.")):
            services.answer(self.conversation, self._inbound("What is the price?"))

        self.assertFalse(Activity.objects.filter(contact=self.contact, kind=Activity.Kind.AUTOMATION).exists())


class TopicDetectionTests(TestCase):
    def test_it_recognises_the_common_business_topics(self):
        self.assertEqual(services.detect_topic("how much does it cost")[0], "pricing")
        self.assertEqual(services.detect_topic("can I get a refund")[0], "refund")
        self.assertEqual(services.detect_topic("is there a warranty")[0], "warranty")
        self.assertEqual(services.detect_topic("when will my delivery arrive")[0], "delivery")

    def test_an_ordinary_message_has_no_topic(self):
        self.assertIsNone(services.detect_topic("hello there"))


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------

class KnowledgeApiTests(KnowledgeFixture):
    def test_listing_bases_creates_the_first_one(self):
        KnowledgeBase.objects.all().delete()

        response = self.client.get("/api/v1/ai/knowledge-bases/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data["results"]), 1)

    def test_a_base_reports_its_document_and_chunk_counts(self):
        self._document()

        response = self.client.get("/api/v1/ai/knowledge-bases/")

        row = next(item for item in response.data["results"] if item["id"] == self.base.pk)
        self.assertEqual(row["document_count"], 1)
        self.assertGreater(row["chunk_total"], 0)

    def test_uploading_a_file_returns_immediately_and_queues_the_work(self):
        with patch("apps.ai.views.process_document_task.delay") as mock_task:
            response = self.client.post(
                "/api/v1/ai/knowledge/upload/",
                {
                    "file": SimpleUploadedFile("prices.pdf", make_pdf(["Premium is 4999."]), content_type="application/pdf"),
                    "knowledge_base": self.base.pk,
                },
                format="multipart",
            )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["status"], KnowledgeDocument.Status.PENDING)
        self.assertEqual(response.data["original_filename"], "prices.pdf")
        mock_task.assert_called_once()

    def test_an_unsupported_file_type_is_refused_before_it_is_stored(self):
        response = self.client.post(
            "/api/v1/ai/knowledge/upload/",
            {"file": SimpleUploadedFile("virus.exe", b"MZ", content_type="application/octet-stream")},
            format="multipart",
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(KnowledgeDocument.objects.filter(source_type="FILE").count(), 0)

    def test_a_file_cannot_be_uploaded_into_another_tenants_base(self):
        other = Tenant.objects.create(name="Rival", slug="kb-api-rival", kind=Tenant.Kind.CUSTOMER)
        foreign = KnowledgeBase.objects.create(tenant=other, name="Theirs")

        response = self.client.post(
            "/api/v1/ai/knowledge/upload/",
            {
                "file": SimpleUploadedFile("prices.pdf", make_pdf(["Premium is 4999."]), content_type="application/pdf"),
                "knowledge_base": foreign.pk,
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, 400)

    def test_reindex_requeues_the_document(self):
        document = self._document()
        with patch("apps.ai.views.process_document_task.delay") as mock_task:
            response = self.client.post(f"/api/v1/ai/knowledge/{document.pk}/reindex/")

        self.assertEqual(response.status_code, 200)
        mock_task.assert_called_once_with(document.pk)

    def test_the_test_panel_answers_without_sending_anything(self):
        self._config()
        self._document()
        before = Message.objects.count()

        with patch.object(services, "complete", return_value=_reply("It is 4999 rupees a month.")):
            response = self.client.post(
                "/api/v1/ai/ask/", {"question": "How much is the premium plan?", "knowledge_base": self.base.pk},
                format="json",
            )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(response.data["answered"])
        self.assertEqual(response.data["confidence"], "GOOD")
        self.assertTrue(response.data["sources"])
        self.assertEqual(Message.objects.count(), before)

    def test_the_test_panel_reports_when_nothing_matches(self):
        self._config()
        self._document()

        with patch.object(services, "complete") as mock_model:
            response = self.client.post(
                "/api/v1/ai/ask/", {"question": "Do you sell helicopters?"}, format="json",
            )

        mock_model.assert_not_called()
        self.assertEqual(response.data["confidence"], "NONE")
        self.assertFalse(response.data["answered"])
        self.assertIn("fallback", response.data["detail"])

    def test_a_test_question_cannot_reach_another_tenants_base(self):
        other = Tenant.objects.create(name="Rival", slug="kb-ask-rival", kind=Tenant.Kind.CUSTOMER)
        foreign = KnowledgeBase.objects.create(tenant=other, name="Theirs")

        response = self.client.post(
            "/api/v1/ai/ask/", {"question": "anything", "knowledge_base": foreign.pk}, format="json",
        )

        self.assertEqual(response.status_code, 400)

    def test_search_reports_the_citation_and_confidence(self):
        self._document()

        response = self.client.post(
            "/api/v1/ai/knowledge/search/", {"query": "premium plan pricing"}, format="json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["confidence"], "GOOD")
        self.assertTrue(response.data["results"][0]["citation"])

    def test_the_answer_log_is_listable_and_tenant_scoped(self):
        other = Tenant.objects.create(name="Rival", slug="kb-log-rival", kind=Tenant.Kind.CUSTOMER)
        KnowledgeAnswer.objects.create(tenant=other, question="theirs", confidence="GOOD")
        KnowledgeAnswer.objects.create(tenant=self.tenant, question="mine", confidence="GOOD")

        response = self.client.get("/api/v1/ai/answers/")

        self.assertEqual([row["question"] for row in response.data["results"]], ["mine"])

    def test_the_gaps_endpoint_lists_unanswered_questions(self):
        KnowledgeAnswer.objects.create(
            tenant=self.tenant, question="Do you deliver to Nepal?", confidence="NONE", is_test=False,
        )
        KnowledgeAnswer.objects.create(
            tenant=self.tenant, question="What is the price?", confidence="GOOD", is_test=False,
        )

        response = self.client.get("/api/v1/ai/answers/gaps/")

        self.assertEqual([row["question"] for row in response.data["results"]], ["Do you deliver to Nepal?"])

    def test_a_base_cannot_be_read_across_tenants(self):
        other = Tenant.objects.create(name="Rival", slug="kb-read-rival", kind=Tenant.Kind.CUSTOMER)
        foreign = KnowledgeBase.objects.create(tenant=other, name="Theirs")

        self.assertEqual(self.client.get(f"/api/v1/ai/knowledge-bases/{foreign.pk}/").status_code, 404)

    def test_deleting_a_base_removes_its_documents_and_chunks(self):
        document = self._document()

        self.client.delete(f"/api/v1/ai/knowledge-bases/{self.base.pk}/")

        self.assertFalse(KnowledgeDocument.objects.filter(pk=document.pk).exists())
        self.assertFalse(KnowledgeChunk.objects.filter(document_id=document.pk).exists())

    def test_everything_needs_authentication(self):
        anonymous = APIClient()
        for url in (
            "/api/v1/ai/knowledge-bases/", "/api/v1/ai/knowledge/", "/api/v1/ai/answers/",
            "/api/v1/ai/knowledge-recommendations/",
        ):
            self.assertEqual(anonymous.get(url).status_code, 401, url)
        self.assertEqual(anonymous.post("/api/v1/ai/ask/", {"question": "x"}, format="json").status_code, 401)


class KnowledgeRecommendationTests(KnowledgeFixture):
    def test_an_empty_knowledge_base_suggests_nothing(self):
        response = self.client.get("/api/v1/ai/knowledge-recommendations/")
        self.assertEqual(response.data["results"], [])

    def test_indexed_documents_with_the_agent_off_suggest_turning_it_on(self):
        self._document()

        response = self.client.get("/api/v1/ai/knowledge-recommendations/")

        self.assertIn("enable_auto_reply", [item["key"] for item in response.data["results"]])

    def test_a_warranty_document_suggests_answering_warranty_questions(self):
        self._document(content=WARRANTY, title="Warranty")

        response = self.client.get("/api/v1/ai/knowledge-recommendations/")

        self.assertIn("topic_warranty", [item["key"] for item in response.data["results"]])

    def test_recommendations_never_reach_another_tenants_documents(self):
        other = Tenant.objects.create(name="Rival", slug="kb-rec-rival", kind=Tenant.Kind.CUSTOMER)
        other_user = get_user_model().objects.create_user(username="kb-rec-owner", password="x", tenant=other)
        other_base = KnowledgeBase.objects.create(tenant=other, name="Theirs")
        KnowledgeDocument.objects.create(
            tenant=other, knowledge_base=other_base, created_by=other_user, title="Warranty",
            content=WARRANTY, status=KnowledgeDocument.Status.READY,
        )

        response = self.client.get("/api/v1/ai/knowledge-recommendations/")

        self.assertEqual(response.data["results"], [])


class KnowledgePresetTests(KnowledgeFixture):
    def test_activating_the_preset_points_the_agent_at_the_base_and_turns_it_on(self):
        self._document()

        response = self.client.post("/api/v1/automation/presets/knowledge_auto_reply/activate/", {
            "knowledge_base": self.base.pk,
            "reply_scope": "EVERY",
            "fallback_action": "SAY_UNKNOWN",
        }, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        config = AIAgentConfig.objects.get(tenant=self.tenant)
        self.assertTrue(config.is_enabled)
        self.assertEqual(config.knowledge_base_id, self.base.pk)
        self.assertEqual(config.reply_scope, "EVERY")
        self.assertEqual(config.fallback_action, "SAY_UNKNOWN")
        self.assertTrue(config.answer_from_knowledge_only)

    def test_a_base_with_nothing_readable_is_refused(self):
        """Switching on an agent with an empty base means every customer gets the fallback."""
        response = self.client.post("/api/v1/automation/presets/knowledge_auto_reply/activate/", {
            "knowledge_base": self.base.pk,
        }, format="json")

        self.assertEqual(response.status_code, 400)
        self.assertFalse(AIAgentConfig.objects.filter(tenant=self.tenant, is_enabled=True).exists())

    def test_another_tenants_base_cannot_be_activated(self):
        other = Tenant.objects.create(name="Rival", slug="kb-preset-rival", kind=Tenant.Kind.CUSTOMER)
        foreign = KnowledgeBase.objects.create(tenant=other, name="Theirs")

        response = self.client.post("/api/v1/automation/presets/knowledge_auto_reply/activate/", {
            "knowledge_base": foreign.pk,
        }, format="json")

        self.assertEqual(response.status_code, 400)

    def test_the_preset_needs_no_whatsapp_number(self):
        """It configures the tenant's agent, not one number's rule."""
        self._document()

        response = self.client.post("/api/v1/automation/presets/knowledge_auto_reply/activate/", {
            "knowledge_base": self.base.pk,
        }, format="json")

        self.assertEqual(response.status_code, 201, response.data)
